# Turn the demo .docx into HTML shaped like Google Docs' "export?format=html":
# classes + <style>, page header first, flat lists with the level in the class
# (lst-kix_<id>-<level>), links wrapped in google.com/url, images on googleusercontent.
import sys, html
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.table import Table

d = Document(sys.argv[1]); out = []; styles = {}; n = [0]; img_n = [0]
numbering = d.part.numbering_part.element
abstract_of = {num.numId: num.abstractNumId.val for num in numbering.num_lst}
fmt = {}
for a in numbering.findall(qn('w:abstractNum')):
    for l in a.findall(qn('w:lvl')):
        fmt[(int(a.get(qn('w:abstractNumId'))), int(l.get(qn('w:ilvl'))))] = l.find(qn('w:numFmt')).get(qn('w:val'))


def cls(color, strike, bold, italic):
    key = (color, strike, bold, italic)
    if key not in styles:
        n[0] += 1; css = []
        if color: css.append('color:#' + color)
        if strike: css.append('text-decoration:line-through')
        css.append('font-weight:700' if bold else 'font-weight:400')
        if italic: css.append('font-style:italic')
        styles[key] = ('c%d' % n[0], ';'.join(css))
    return styles[key][0]


def runs(p):
    s = ''
    for el in p._p:
        if el.tag == qn('w:r'):
            r = Run(el, p)
            if el.findall('.//' + qn('w:drawing')):
                img_n[0] += 1
                s += '<span class="imx"><img src="https://lh7-rt.googleusercontent.com/docsz/demo-image-%d" style="width:300px"></span>' % img_n[0]
                continue
            col = str(r.font.color.rgb) if r.font.color and r.font.color.type else None
            s += '<span class="%s">%s</span>' % (cls(col, bool(r.font.strike), bool(r.bold), bool(r.italic)), html.escape(r.text))
        elif el.tag == qn('w:hyperlink'):
            rid = el.get(qn('r:id')); url = p.part.rels[rid].target_ref
            for rr in el.findall(qn('w:r')):
                t = ''.join(x.text for x in rr.findall(qn('w:t')))
                c = rr.find(qn('w:rPr')); col = None; st = False
                if c is not None:
                    cc = c.find(qn('w:color')); col = cc.get(qn('w:val')) if cc is not None else None; st = c.find(qn('w:strike')) is not None
                s += '<span class="%s"><a class="c99" href="https://www.google.com/url?q=%s&amp;sa=D&amp;source=editors">%s</a></span>' % (cls(col, st, False, False), html.escape(url), html.escape(t))
    return s


def list_info(p):
    pPr = p._p.pPr
    if pPr is None or pPr.numPr is None or pPr.numPr.numId is None:
        return None
    nid = pPr.numPr.numId.val; lvl = pPr.numPr.ilvl.val if pPr.numPr.ilvl is not None else 0
    kind = 'ol' if fmt.get((abstract_of.get(nid), lvl)) == 'decimal' else 'ul'
    return (nid, lvl, kind)


# page header first, like Google does
hdr = d.sections[0].header
out.append('<div>' + ''.join('<p class="pb2"><span class="%s">%s</span></p>' % (cls('666666', False, False, False), html.escape(p.text)) for p in hdr.paragraphs) + '</div>')
cur = None   # (nid, lvl, kind) of the open list
for el in d.element.body:
    if el.tag == qn('w:p'):
        p = Paragraph(el, d); st = p.style.name; li = list_info(p)
        if li != cur:
            if cur: out.append('</%s>' % cur[2])
            if li: out.append('<%s class="lst4 lst-kix_demo%s-%d start">' % (li[2], li[0], li[1]))
            cur = li
        if li: out.append('<li class="lbx li-bullet-0">%s</li>' % runs(p)); continue
        tag = {'Heading 1': 'h1', 'Heading 2': 'h2', 'Heading 3': 'h3'}.get(st, 'p')
        out.append('<%s class="pb2" id="h.x%d">%s</%s>' % (tag, len(out), runs(p), tag) if tag != 'p' else '<p class="pb2">%s</p>' % runs(p))
    elif el.tag == qn('w:tbl'):
        if cur: out.append('</%s>' % cur[2]); cur = None
        t = Table(el, d); out.append('<table class="tbx"><tbody>')
        for row in t.rows:
            out.append('<tr class="trx">' + ''.join('<td class="tc7" colspan="1" rowspan="1">' + ''.join('<p class="pb2">%s</p>' % runs(pp) for pp in c.paragraphs) + '</td>' for c in row.cells) + '</tr>')
        out.append('</tbody></table>')
if cur: out.append('</%s>' % cur[2])
css = ''.join('.%s{%s}' % (v[0], v[1]) for v in styles.values()) + '.c99{color:#1155cc;text-decoration:underline}.lst-kix_demo-0>li:before{content:"x"}'
print('<html><head><meta content="text/html; charset=UTF-8" http-equiv="content-type"><style type="text/css">' + css + '</style></head><body class="bdx doc-content">' + ''.join(out) + '</body></html>')
